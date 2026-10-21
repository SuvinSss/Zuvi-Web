"""Downloadable Excel template: Products sheet, Instructions and Reference."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from catalog.models import Brand, ProductCategory, ProductUnit, Tag
from stores.models import Store

from .columns import COLUMNS
from .limits import limit

REQUIRED_FILL = PatternFill("solid", fgColor="14532D")
OPTIONAL_FILL = PatternFill("solid", fgColor="647067")


def build_template() -> bytes:
    wb = Workbook()
    products = wb.active
    products.title = "Products"
    for index, column in enumerate(COLUMNS, start=1):
        cell = products.cell(row=1, column=index, value=column.key)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = REQUIRED_FILL if column.required else OPTIONAL_FILL
        cell.comment = Comment(("Required. " if column.required else "Optional. ") + column.description, "Zuvi")
        products.column_dimensions[get_column_letter(index)].width = max(16, len(column.key) + 4)
    products.freeze_panes = "A2"
    unit_col = get_column_letter([c.key for c in COLUMNS].index("unit") + 1)
    validation = DataValidation(type="list", formula1='"' + ",".join(ProductUnit.values) + '"', allow_blank=True)
    products.add_data_validation(validation)
    validation.add(f"{unit_col}2:{unit_col}{limit('MAX_ROWS') + 1}")

    help_sheet = wb.create_sheet("Instructions")
    help_sheet.append(["Zuvi bulk product import"])
    help_sheet["A1"].font = Font(bold=True, size=14)
    for line in (
        "Fill the Products sheet, one product per row. Dark green headers are required; grey are optional.",
        f"Limits: {limit('MAX_ROWS')} rows and {limit('MAX_FILE_BYTES') // (1024 * 1024)} MB per file. Formulas are not allowed; enter values.",
        "Store + SKU identify a product. Import modes decide what happens when it already exists:",
        "  create: existing SKUs are reported as errors.  skip: existing SKUs are ignored.  update: existing SKUs are updated.",
        "Existing products are never changed silently: updates only happen in update mode and are counted in the result.",
        "Stock is not imported: stock changes must go through inventory so every movement is recorded.",
        "Margin, discount and final price are set by management after import. Products start as Draft unless you choose Pending.",
        "Images: either list file names in image_file and upload a ZIP of images, or tick 'Allow image URLs' and use image_url.",
        "Products without an image automatically show the Zuvi category/brand placeholder.",
        "", "Columns",
    ):
        help_sheet.append([line])
    help_sheet.append(["column", "required", "description", "example"])
    for cell in help_sheet[help_sheet.max_row]:
        cell.font = Font(bold=True)
    for column in COLUMNS:
        help_sheet.append([column.key, "required" if column.required else "optional", column.description, column.example])
    help_sheet.column_dimensions["A"].width = 24
    help_sheet.column_dimensions["C"].width = 90
    for row in help_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=False, vertical="top")

    reference = wb.create_sheet("Reference")
    reference.append(["store_code", "store_name", "category", "brand", "tag", "unit"])
    for cell in reference[1]:
        cell.font = Font(bold=True)
    stores = list(Store.objects.order_by("name").values_list("store_code", "name")[:500])
    categories = [str(c) for c in ProductCategory.objects.filter(is_active=True).select_related("parent__parent")][:1000]
    brands = list(Brand.objects.filter(is_active=True).values_list("name", flat=True)[:1000])
    tags = list(Tag.objects.filter(is_active=True).values_list("name", flat=True)[:1000])
    units = list(ProductUnit.values)
    for i in range(max(len(stores), len(categories), len(brands), len(tags), len(units))):
        store = stores[i] if i < len(stores) else ("", "")
        reference.append([
            store[0], store[1],
            categories[i] if i < len(categories) else "",
            brands[i] if i < len(brands) else "",
            tags[i] if i < len(tags) else "",
            units[i] if i < len(units) else "",
        ])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
