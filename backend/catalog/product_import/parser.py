"""Read .xlsx / .csv into plain row dictionaries, defensively.

Nothing here trusts the file: the xlsx container is size-checked before it is
opened (decompression bombs), XML is parsed with defusedxml (installed, and
asserted below), formula cells are refused and text values that look like
formulas are rejected.
"""
import csv
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

import openpyxl
from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from .columns import KEYS, REQUIRED, TEXT_KEYS, resolve_headers
from .limits import limit

assert openpyxl.DEFUSEDXML, "defusedxml must be installed so openpyxl parses XML safely"

FORMULA_PREFIXES = ("=", "@")


@dataclass
class RawRow:
    number: int  # spreadsheet row number (header is row 1)
    values: dict = field(default_factory=dict)
    errors: list = field(default_factory=list)  # [(column, message)]


def _cell_value(cell, key, errors):
    if getattr(cell, "data_type", "") == "f":
        errors.append((key, "Formulas are not allowed. Enter the value instead."))
        return None
    value = cell.value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return Decimal(repr(value))
    if isinstance(value, int):
        return Decimal(value)
    return _clean_text(value, key, errors)


def _clean_text(value, key, errors):
    if value is None:
        return None
    if isinstance(value, (date, Decimal)):
        return value
    text = str(value).strip()
    if not text:
        return None
    if key in TEXT_KEYS and text.startswith(FORMULA_PREFIXES):
        errors.append((key, "Text may not start with '=' or '@'."))
        return None
    return text


def _check_container(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > limit("MAX_UNCOMPRESSED_BYTES"):
                raise ValidationError("The workbook expands to an unreasonable size.")
    except zipfile.BadZipFile:
        raise ValidationError("This is not a valid .xlsx workbook.") from None


class MappingNeeded(ValidationError):
    """Required columns could not be identified. ``info`` feeds the manual mapping screen."""

    def __init__(self, message, info):
        super().__init__(message)
        self.info = info


def _prepare(raw_headers, mapping, defaults):
    """Resolve headers; raise MappingNeeded when a required column cannot be identified."""
    resolved, proposal = resolve_headers(raw_headers, mapping)
    present = set(resolved.values()) | {k for k, v in (defaults or {}).items() if v}
    missing = [key for key in REQUIRED if key not in present]
    info = {"headers": [str(h) for h in raw_headers if str(h or "").strip()], "proposal": proposal, "missing_required": missing}
    if missing:
        raise MappingNeeded(
            "Missing required column(s): " + ", ".join(missing) + ". Download the template for the exact headers.", info
        )
    return resolved, info


def _finish(body, info):
    ignored = [p["source"] for p in info["proposal"] if not p["key"]]
    return body, ignored


def inspect_headers(data: bytes, filename: str, mapping=None, defaults=None):
    """Header analysis only (no rows). Never raises for unmapped columns; used by the mapping screen."""
    try:
        return parse_workbook(data, filename, mapping=mapping, defaults=defaults, headers_only=True)
    except MappingNeeded as exc:
        return exc.info


def parse_workbook(data: bytes, filename: str, mapping=None, defaults=None, headers_only=False):
    """Return ``(rows, ignored_columns)``; raises ValidationError for unusable files.

    ``mapping`` maps source headers to canonical keys (user decisions); ``defaults`` supplies a value for a
    required column the file does not have (e.g. ``store_code`` chosen on screen).
    """
    if len(data) > limit("MAX_FILE_BYTES"):
        raise ValidationError(f"The file must be {limit('MAX_FILE_BYTES') // (1024 * 1024)} MB or smaller.")
    if not data:
        raise ValidationError("The file is empty.")
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _parse_csv(data, mapping, defaults, headers_only)
    if name.endswith(".xlsx"):
        return _parse_xlsx(data, mapping, defaults, headers_only)
    raise ValidationError("Upload an .xlsx or .csv file.")


def _parse_xlsx(data, mapping=None, defaults=None, headers_only=False):
    _check_container(data)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception:  # openpyxl raises many types for corrupt files
        raise ValidationError("This is not a valid .xlsx workbook.") from None
    try:
        sheet = wb["Products"] if "Products" in wb.sheetnames else wb.worksheets[0]
        iterator = sheet.iter_rows()
        try:
            header_cells = next(iterator)
        except StopIteration:
            raise ValidationError("The workbook has no header row.") from None
        header, info = _prepare([c.value for c in header_cells], mapping, defaults)
        if headers_only:
            return info
        body = []
        for number, cells in enumerate(iterator, start=2):
            if len(body) >= limit("MAX_ROWS"):
                raise ValidationError(f"The file has more than {limit('MAX_ROWS')} rows. Split it into smaller files.")
            row = RawRow(number)
            empty = True
            for index, cell in enumerate(cells):
                key = header.get(index)
                if key is None:
                    continue
                value = _cell_value(cell, key, row.errors)
                if value is not None or getattr(cell, "data_type", "") == "f":
                    empty = False
                row.values[key] = value
            if not empty:
                body.append(row)
        return _finish(body, info)
    finally:
        wb.close()


def _parse_csv(data, mapping=None, defaults=None, headers_only=False):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ValidationError("CSV files must be UTF-8 encoded.") from None
    reader = csv.reader(io.StringIO(text))
    try:
        header_row = next(reader)
    except StopIteration:
        raise ValidationError("The file has no header row.") from None
    header, info = _prepare(header_row, mapping, defaults)
    if headers_only:
        return info
    body = []
    for number, cells in enumerate(reader, start=2):
        if not any((c or "").strip() for c in cells):
            continue
        if len(body) >= limit("MAX_ROWS"):
            raise ValidationError(f"The file has more than {limit('MAX_ROWS')} rows. Split it into smaller files.")
        row = RawRow(number)
        for index, value in enumerate(cells):
            key = header.get(index)
            if key is not None:
                row.values[key] = _clean_text(value, key, row.errors)
        body.append(row)
    return _finish(body, info)
