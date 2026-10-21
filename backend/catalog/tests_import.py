"""Bulk import: template, parsing, validation, modes, images, atomicity, security."""
import io
import zipfile
from datetime import date, timedelta
from decimal import Decimal
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .models import Product, ProductImage, ProductImportBatch, ProductStatus
from .product_import import images as import_images
from .product_import import reports
from .product_import.columns import KEYS
from .product_import.parser import parse_workbook
from .tests import CatalogTestMixin

User = get_user_model()
PASSWORD = "secure-password-123"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_bytes(rows, header=KEYS):
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(list(header))
    for row in rows:
        ws.append([row.get(key) for key in header])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def zip_bytes(files):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return out.getvalue()


class ImportBase(CatalogTestMixin, TestCase):
    def setUp(self):
        cache.clear()
        self.store = self.create_store()
        self.category = self.create_product_category(name="Groceries")
        self.brand = self.create_brand(name="Zuvi Farms")
        self.tag = self.create_tag(name="organic")
        self.user = User.objects.create_superuser("root", "root@example.com", PASSWORD)
        self.client = Client(enforce_csrf_checks=True)
        self.client.force_login(self.user)
        self.client.get(reverse("api_auth_session"))
        self.token = self.client.cookies["csrftoken"].value

    def row(self, **overrides):
        data = {"store_code": self.store.store_code, "sku": "RICE-1", "name": "Rice", "category": "Groceries",
                "store_price": "99.50"}
        data.update(overrides)
        return data

    def preview(self, rows, *, mode="create", zip_data=None, name="products.xlsx", content=None, **fields):
        files = {"file": SimpleUploadedFile(name, content or workbook_bytes(rows), content_type=XLSX)}
        if zip_data is not None:
            files["images_zip"] = SimpleUploadedFile("images.zip", zip_data, content_type="application/zip")
        return self.client.post(reverse("api_import_preview"), {**files, "mode": mode, **fields}, HTTP_X_CSRFTOKEN=self.token)

    def commit(self, batch_id):
        return self.client.post(reverse("api_import_commit", args=[batch_id]), HTTP_X_CSRFTOKEN=self.token)


class TemplateTests(ImportBase):
    def test_template_has_marked_columns_instructions_and_reference_data(self):
        response = self.client.get(reverse("api_import_template"))
        self.assertEqual(response["Content-Type"], XLSX)
        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["Products", "Instructions", "Reference"])
        self.assertEqual([c.value for c in wb["Products"][1]], list(KEYS))
        required = [c.value for c in wb["Products"][1] if c.fill.fgColor.rgb.endswith("14532D")]
        self.assertEqual(required, ["store_code", "sku", "name", "category", "store_price"])
        reference = [[c.value for c in row] for row in wb["Reference"].iter_rows(min_row=2)]
        self.assertIn(self.store.store_code, [r[0] for r in reference])

    def test_template_round_trips_through_the_parser(self):
        wb = load_workbook(io.BytesIO(self.client.get(reverse("api_import_template")).content))
        out = io.BytesIO()
        wb.save(out)
        rows, ignored = parse_workbook(out.getvalue(), "t.xlsx")
        self.assertEqual((rows, ignored), ([], []))


class ParserSecurityTests(ImportBase):
    def test_formula_cells_are_rejected(self):
        wb = Workbook()
        ws = wb.active
        ws.append(list(KEYS))
        ws.append([self.store.store_code, "F1", "=1+1", "Groceries", 5])
        out = io.BytesIO()
        wb.save(out)
        rows, _ = parse_workbook(out.getvalue(), "f.xlsx")
        self.assertTrue(any("Formulas are not allowed" in m for _, m in rows[0].errors))

    def test_text_starting_with_equals_or_at_is_rejected(self):
        body = ",".join(KEYS) + "\n" + f"{self.store.store_code},S1,@SUM(A1),Groceries,5\n"
        rows, _ = parse_workbook(body.encode(), "x.csv")
        self.assertTrue(rows[0].errors)

    def test_non_xlsx_and_missing_columns_and_row_limits(self):
        with self.assertRaises(ValidationError):
            parse_workbook(b"PK not really", "x.xlsx")
        with self.assertRaises(ValidationError):
            parse_workbook(b"a,b\n1,2\n", "x.csv")
        with self.assertRaises(ValidationError):
            parse_workbook(b"x", "x.txt")
        with self.settings(PRODUCT_IMPORT_MAX_ROWS=2), self.assertRaisesMessage(ValidationError, "more than 2 rows"):
            parse_workbook(workbook_bytes([self.row(sku=f"S{i}") for i in range(3)]), "x.xlsx")

    def test_oversized_file_and_zip_bomb_workbook(self):
        with self.settings(PRODUCT_IMPORT_MAX_FILE_BYTES=10), self.assertRaises(ValidationError):
            parse_workbook(workbook_bytes([self.row()]), "x.xlsx")
        with self.settings(PRODUCT_IMPORT_MAX_UNCOMPRESSED_BYTES=100), self.assertRaisesMessage(ValidationError, "unreasonable"):
            parse_workbook(workbook_bytes([self.row()]), "x.xlsx")

    def test_error_report_neutralises_formula_injection(self):
        batch = ProductImportBatch(errors=[{"row": 2, "sku": "=cmd()", "column": "sku", "message": "@evil"}])
        batch.pk = "00000000-0000-0000-0000-000000000001"
        text = reports.csv_response(batch).content.decode()
        self.assertIn("'=cmd()", text)
        self.assertIn("'@evil", text)


class ZipAndUrlSafetyTests(TestCase):
    def test_zip_path_traversal_and_bad_types_are_rejected(self):
        for name in ("../evil.jpg", "/abs.jpg", "a/../../evil.png"):
            with self.assertRaises(ValidationError, msg=name):
                import_images.ZipImages(zip_bytes({name: b"x"}))
        with self.assertRaisesMessage(ValidationError, "Unsupported file"):
            import_images.ZipImages(zip_bytes({"run.sh": b"#!/bin/sh"}))
        with self.assertRaisesMessage(ValidationError, "Duplicate"):
            import_images.ZipImages(zip_bytes({"a/x.jpg": b"1", "b/x.jpg": b"2"}))
        with self.assertRaises(ValidationError):
            import_images.ZipImages(b"not a zip")

    def test_zip_limits_and_ignored_metadata_files(self):
        with self.settings(PRODUCT_IMPORT_ZIP_MAX_FILES=1), self.assertRaises(ValidationError):
            import_images.ZipImages(zip_bytes({"a.jpg": b"1", "b.jpg": b"2"}))
        with self.settings(PRODUCT_IMAGE_MAX_BYTES=5), self.assertRaises(ValidationError):
            import_images.ZipImages(zip_bytes({"a.jpg": b"123456"}))
        ok = import_images.ZipImages(zip_bytes({"a.jpg": b"1", "__MACOSX/._a.jpg": b"x", ".DS_Store": b"x"}))
        self.assertIn("A.JPG", ok and ["A.JPG"]) if "a.jpg" in ok else self.fail("a.jpg missing")

    def test_symlinks_are_rejected(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as archive:
            info = zipfile.ZipInfo("link.jpg")
            info.external_attr = (0o120777 << 16)
            archive.writestr(info, "/etc/passwd")
        with self.assertRaisesMessage(ValidationError, "Symbolic"):
            import_images.ZipImages(out.getvalue())

    def test_url_ssrf_blocks_private_loopback_and_link_local(self):
        for url in ("http://127.0.0.1/x.jpg", "http://10.0.0.5/x.jpg", "http://169.254.169.254/latest/meta-data",
                    "http://[::1]/x.jpg", "http://192.168.1.1/x.jpg", "http://0.0.0.0/x.jpg"):
            with self.assertRaises(ValidationError, msg=url):
                import_images.validate_image_url(url)
        for url in ("file:///etc/passwd", "ftp://example.com/x.jpg", "http://user:pw@example.com/x.jpg", "javascript:1"):
            with self.assertRaises(ValidationError, msg=url):
                import_images.validate_image_url(url)

    def test_hostname_resolving_to_private_address_is_blocked(self):
        with mock.patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("10.1.2.3", 80))]):
            with self.assertRaises(ValidationError):
                import_images.validate_image_url("http://innocent.example.com/x.jpg")
        with mock.patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("93.184.216.34", 80)), (2, 1, 6, "", ("127.0.0.1", 80))]):
            with self.assertRaises(ValidationError):  # one private answer is enough to refuse
                import_images.validate_image_url("http://rebind.example.com/x.jpg")

    def test_fetch_fails_closed_when_peer_cannot_be_verified_or_is_private(self):
        response = mock.MagicMock(status_code=200, headers={})
        response.__enter__.return_value = response
        with mock.patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("93.184.216.34", 80))]), \
                mock.patch("requests.get", return_value=response):
            response.raw._connection.sock.getpeername.return_value = ("127.0.0.1", 80)  # DNS rebinding
            with self.assertRaises(ValidationError):
                import_images.fetch_image_url("http://example.com/x.jpg")
            response.raw = None  # cannot determine the peer -> refuse
            with self.assertRaises(ValidationError):
                import_images.fetch_image_url("http://example.com/x.jpg")

    def test_fetch_enforces_size_cap_and_refuses_redirects(self):
        response = mock.MagicMock(status_code=200, headers={})
        response.__enter__.return_value = response
        response.raw._connection.sock.getpeername.return_value = ("93.184.216.34", 80)
        response.iter_content.return_value = [b"x" * 600, b"x" * 600]
        with self.settings(PRODUCT_IMAGE_MAX_BYTES=1000), \
                mock.patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("93.184.216.34", 80))]), \
                mock.patch("requests.get", return_value=response) as get:
            with self.assertRaisesMessage(ValidationError, "too large"):
                import_images.fetch_image_url("http://example.com/x.jpg")
            self.assertFalse(get.call_args.kwargs["allow_redirects"])
        redirect = mock.MagicMock(status_code=302, headers={})
        redirect.__enter__.return_value = redirect
        with mock.patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("93.184.216.34", 80))]), \
                mock.patch("requests.get", return_value=redirect):
            with self.assertRaisesMessage(ValidationError, "HTTP 302"):
                import_images.fetch_image_url("http://example.com/x.jpg")


class PreviewTests(ImportBase):
    def test_preview_reports_counts_row_column_errors_and_writes_no_products(self):
        rows = [
            self.row(),
            self.row(sku="BAD-1", store_price="-5", category="Nope"),
            self.row(sku="RICE-1"),  # duplicate of row 2 in the same file
            self.row(sku="BAD-2", store_code="NOPE", manufacturing_date=date.today() + timedelta(days=3)),
        ]
        response = self.preview(rows)
        self.assertEqual(response.status_code, 201, response.content)
        data = response.json()
        self.assertEqual(data["counts"]["total"], 4)
        self.assertEqual(data["counts"]["valid"], 1)
        self.assertEqual(data["counts"]["invalid"], 3)
        self.assertEqual(data["counts"]["duplicates"], 1)
        errors = {(e["row"], e["column"]) for e in data["errors"]}
        self.assertTrue({(3, "store_price"), (3, "category"), (4, "sku"), (5, "store_code"), (5, "manufacturing_date")} <= errors)
        self.assertEqual(Product.objects.count(), 0)
        self.assertEqual(data["rows"][0]["status"], "valid")

    def test_existing_sku_follows_the_selected_mode(self):
        self.create_product(store=self.store, category=self.category, sku="RICE-1", name="Old")
        create = self.preview([self.row()], mode="create").json()
        self.assertEqual(create["counts"]["invalid"], 1)
        self.assertEqual(create["counts"]["duplicates"], 1)
        skip = self.preview([self.row()], mode="skip").json()
        self.assertEqual((skip["counts"]["invalid"], skip["counts"]["skip"]), (0, 1))
        update = self.preview([self.row()], mode="update").json()
        self.assertEqual((update["counts"]["invalid"], update["counts"]["update"]), (0, 1))

    def test_unknown_brand_unit_tag_and_bad_numbers(self):
        data = self.preview([self.row(brand="Ghost", unit="BUCKET", tags="organic, ghost", store_price="1.234",
                                      unit_value="abc", expiry_date="31/12/2027")]).json()
        columns = {e["column"] for e in data["errors"]}
        self.assertEqual(columns, {"brand", "unit", "tags", "store_price", "unit_value", "expiry_date"})

    def test_invalid_files_are_400s_not_500s(self):
        self.assertEqual(self.preview([], content=b"not a workbook").status_code, 400)
        self.assertEqual(self.preview([], name="x.pdf", content=b"%PDF").status_code, 400)
        self.assertEqual(self.client.post(reverse("api_import_preview"), {}, HTTP_X_CSRFTOKEN=self.token).status_code, 400)

    def test_image_references_are_checked_against_the_zip(self):
        zip_data = zip_bytes({"rice.jpg": self._jpeg_bytes()})
        data = self.preview([self.row(image_file="rice.jpg"), self.row(sku="R2", image_file="missing.jpg")], zip_data=zip_data).json()
        self.assertEqual(data["rows"][0]["image"], "zip")
        self.assertEqual([e["column"] for e in data["errors"]], ["image_file"])

    def test_image_urls_are_off_unless_explicitly_enabled(self):
        row = self.row(image_url="https://example.com/a.jpg")
        data = self.preview([row]).json()
        self.assertEqual(data["errors"][0]["column"], "image_url")
        with mock.patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("93.184.216.34", 443))]):
            ok = self.preview([row], allow_image_urls="true").json()
        self.assertEqual(ok["counts"]["invalid"], 0)
        private = self.preview([self.row(image_url="http://127.0.0.1/x.png")], allow_image_urls="true").json()
        self.assertEqual(private["errors"][0]["column"], "image_url")

    def test_numeric_sku_cells_stay_text(self):
        wb = Workbook()
        ws = wb.active
        ws.append(list(KEYS))
        ws.append([self.store.store_code, 12345, "Num SKU", "Groceries", 10.5])
        out = io.BytesIO()
        wb.save(out)
        data = self.preview([], content=out.getvalue()).json()
        self.assertEqual(data["rows"][0]["sku"], "12345")
        self.assertEqual(data["counts"]["invalid"], 0)

    def test_csv_is_supported(self):
        body = ",".join(KEYS) + f"\n{self.store.store_code},C1,CSV Rice,Groceries,12.00\n"
        data = self.preview([], name="p.csv", content=body.encode()).json()
        self.assertEqual(data["counts"]["valid"], 1)


class CommitTests(ImportBase):
    def test_commit_creates_products_via_services_with_history_and_defaults(self):
        batch = self.preview([self.row(brand="Zuvi Farms", tags="organic", unit="KG", unit_value="5",
                                       description="Good rice")]).json()
        response = self.commit(batch["id"])
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(body["status"], "COMMITTED")
        self.assertEqual(body["summary"]["created"], 1)
        self.assertEqual(body["summary"]["outcome"], "complete")
        self.assertEqual(body["summary"]["fallback_images"], 1)
        product = Product.objects.get()
        self.assertEqual(product.status, ProductStatus.DRAFT)
        self.assertEqual(product.store_price, Decimal("99.50"))
        self.assertEqual(product.stock_quantity, Decimal("0"))
        self.assertIsNone(product.final_price)
        self.assertEqual(product.brand, self.brand)
        self.assertEqual(list(product.tags.all()), [self.tag])
        self.assertEqual(product.created_by, self.user)
        self.assertEqual(product.status_history.count(), 1)

    def test_pending_start_status_is_supported_and_validated(self):
        batch = self.preview([self.row()], initial_status="PENDING").json()
        self.commit(batch["id"])
        self.assertEqual(Product.objects.get().status, ProductStatus.PENDING)
        self.assertEqual(self.preview([self.row(sku="Z")], initial_status="APPROVED").status_code, 400)

    def test_double_submit_is_idempotent(self):
        batch = self.preview([self.row()]).json()
        first, second = self.commit(batch["id"]), self.commit(batch["id"])
        self.assertEqual((first.status_code, second.status_code), (200, 200))
        self.assertEqual(Product.objects.count(), 1)

    def test_same_file_cannot_be_imported_twice(self):
        content = workbook_bytes([self.row()])
        first = self.preview([], content=content).json()
        self.commit(first["id"])
        again = self.preview([], content=content, mode="skip").json()
        self.assertTrue(again["already_imported"])
        response = self.commit(again["id"])
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["code"], "already_imported")
        self.assertEqual(Product.objects.count(), 1)

    def test_default_is_all_or_nothing(self):
        batch = self.preview([self.row(), self.row(sku="BAD", store_price="-1")]).json()
        response = self.commit(batch["id"])
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["code"], "invalid_rows")
        self.assertEqual(Product.objects.count(), 0)

    def test_valid_rows_only_reports_partial_not_complete(self):
        batch = self.preview([self.row(), self.row(sku="BAD", store_price="-1")], all_or_nothing="false").json()
        body = self.commit(batch["id"]).json()
        self.assertEqual((body["summary"]["created"], body["summary"]["failed"], body["summary"]["outcome"]),
                         (1, 1, "partial"))
        self.assertEqual(Product.objects.count(), 1)
        errors = self.client.get(reverse("api_import_errors", args=[batch["id"]]), {"type": "csv"}).content.decode()
        self.assertIn("store_price", errors)
        xlsx = self.client.get(reverse("api_import_errors", args=[batch["id"]]))
        self.assertEqual(xlsx["Content-Type"], XLSX)

    def test_a_row_failing_at_save_time_rolls_everything_back_in_strict_mode(self):
        batch = self.preview([self.row(), self.row(sku="R2")]).json()
        real = Product.save
        calls = []

        def flaky(self_, *a, **k):
            calls.append(1)
            if len(calls) == 2:
                raise RuntimeError("boom")
            return real(self_, *a, **k)

        with mock.patch.object(Product, "save", flaky):
            response = self.commit(batch["id"])
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["code"], "rolled_back")
        self.assertEqual(Product.objects.count(), 0)  # first row was rolled back too
        self.assertEqual(ProductImportBatch.objects.get().status, "FAILED")
        self.assertNotIn("boom", response.content.decode())  # internals stay in the log

    def test_a_row_failing_at_save_time_only_skips_that_row_in_partial_mode(self):
        batch = self.preview([self.row(), self.row(sku="R2")], all_or_nothing="false").json()
        real = Product.save
        calls = []

        def flaky(self_, *a, **k):
            calls.append(1)
            if len(calls) == 2:
                raise RuntimeError("boom")
            return real(self_, *a, **k)

        with mock.patch.object(Product, "save", flaky):
            body = self.commit(batch["id"]).json()
        self.assertEqual((body["summary"]["created"], body["summary"]["failed"]), (1, 1))
        self.assertEqual(list(Product.objects.values_list("sku", flat=True)), ["RICE-1"])

    def test_update_mode_updates_reports_and_never_touches_unchanged_rows(self):
        self.create_product(store=self.store, category=self.category, sku="RICE-1", name="Old name",
                            store_price=Decimal("99.50"))
        self.create_product(store=self.store, category=self.category, sku="SAME", name="Same", store_price=Decimal("5.00"))
        rows = [self.row(name="New name"), self.row(sku="SAME", name="Same", store_price="5.00", unit="PIECE"), self.row(sku="NEW")]
        body = self.commit(self.preview(rows, mode="update").json()["id"]).json()
        self.assertEqual((body["summary"]["updated"], body["summary"]["created"], body["summary"]["skipped"]), (1, 1, 1))
        self.assertEqual(Product.objects.get(sku="RICE-1").name, "New name")

    def test_skip_mode_leaves_existing_products_untouched(self):
        existing = self.create_product(store=self.store, category=self.category, sku="RICE-1", name="Keep me")
        body = self.commit(self.preview([self.row(name="Overwrite?"), self.row(sku="N")], mode="skip").json()["id"]).json()
        self.assertEqual((body["summary"]["skipped"], body["summary"]["created"]), (1, 1))
        existing.refresh_from_db()
        self.assertEqual(existing.name, "Keep me")

    def test_data_changed_between_preview_and_commit_is_revalidated(self):
        batch = self.preview([self.row()]).json()
        self.create_product(store=self.store, category=self.category, sku="RICE-1")  # appears after preview
        response = self.commit(batch["id"])
        self.assertEqual(response.status_code, 409)
        self.assertEqual(Product.objects.count(), 1)

    def test_staged_files_are_deleted_after_commit(self):
        from django.core.files.storage import default_storage
        batch = self.preview([self.row()]).json()
        path = ProductImportBatch.objects.get().workbook_path
        self.assertTrue(default_storage.exists(path))
        self.commit(batch["id"])
        self.assertFalse(default_storage.exists(path))
        self.assertEqual(ProductImportBatch.objects.get().workbook_path, "")


class ImportImageTests(ImportBase):
    def test_zip_images_go_through_the_shared_pipeline(self):
        zip_data = zip_bytes({"photos/rice.jpg": self._jpeg_bytes((3000, 2000)), "dal.png": self._png_bytes((50, 50))})
        rows = [self.row(image_file="rice.jpg", image_alt_text="Rice"), self.row(sku="DAL", image_file="dal.png"),
                self.row(sku="NOPIC")]
        body = self.commit(self.preview(rows, zip_data=zip_data).json()["id"]).json()
        self.assertEqual((body["summary"]["images_processed"], body["summary"]["fallback_images"]), (2, 1))
        image = ProductImage.objects.get(product__sku="RICE-1")
        self.assertEqual(image.alt_text, "Rice")
        self.assertTrue(image.thumbnail.name.endswith(".webp"))
        self.assertEqual(max(image.width, image.height), 2400)
        self.assertFalse(Product.objects.get(sku="NOPIC").images.exists())

    def test_identical_images_in_one_import_are_stored_once(self):
        data = self._jpeg_bytes((300, 300))
        zip_data = zip_bytes({"a.jpg": data, "b.jpg": data})
        self.commit(self.preview([self.row(image_file="a.jpg"), self.row(sku="R2", image_file="b.jpg")], zip_data=zip_data).json()["id"])
        names = set(ProductImage.objects.values_list("image", flat=True))
        self.assertEqual(len(names), 1)

    def test_corrupt_image_fails_that_row_without_corrupting_others(self):
        zip_data = zip_bytes({"ok.jpg": self._jpeg_bytes(), "bad.jpg": b"this is not a jpeg"})
        rows = [self.row(image_file="ok.jpg"), self.row(sku="BAD", image_file="bad.jpg")]
        body = self.commit(self.preview(rows, zip_data=zip_data, all_or_nothing="false").json()["id"]).json()
        self.assertEqual((body["summary"]["created"], body["summary"]["failed"]), (1, 1))
        self.assertFalse(Product.objects.filter(sku="BAD").exists())  # no half-created product
        self.assertEqual(ProductImage.objects.count(), 1)
        self.assertEqual(body["errors"][-1]["row"], 3)

    def test_url_images_are_downloaded_only_when_enabled(self):
        row = self.row(image_url="https://example.com/a.jpg")
        with mock.patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("93.184.216.34", 443))]):
            batch = self.preview([row], allow_image_urls="true").json()
        with mock.patch("catalog.product_import.service.fetch_image_url", return_value=self._jpeg_bytes()) as fetch, \
                mock.patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("93.184.216.34", 443))]):
            body = self.commit(batch["id"]).json()
        fetch.assert_called_once_with("https://example.com/a.jpg")
        self.assertEqual(body["summary"]["images_processed"], 1)


class ImportPermissionTests(ImportBase):
    def test_admin_without_add_permission_is_refused(self):
        admin = User.objects.create_user("mgr", "m@example.com", PASSWORD, role="ADMIN", is_staff=True)
        self.grant_catalog_perms(admin, "view_product")
        c = Client()
        c.force_login(admin)
        self.assertEqual(c.get(reverse("api_import_template")).status_code, 403)

    def test_update_mode_and_urls_need_change_permission(self):
        admin = User.objects.create_user("mgr", "m@example.com", PASSWORD, role="ADMIN", is_staff=True)
        self.grant_catalog_perms(admin, "view_product", "add_product")
        c = Client(enforce_csrf_checks=True)
        c.force_login(admin)
        c.get(reverse("api_auth_session"))
        token = c.cookies["csrftoken"].value
        for extra in ({"mode": "update"}, {"allow_image_urls": "true"}):
            file = SimpleUploadedFile("p.xlsx", workbook_bytes([self.row()]), content_type=XLSX)
            response = c.post(reverse("api_import_preview"), {"file": file, **extra}, HTTP_X_CSRFTOKEN=token)
            self.assertEqual(response.status_code, 403, extra)

    def test_users_cannot_see_each_others_imports(self):
        batch = self.preview([self.row()]).json()
        other = User.objects.create_user("mgr", "m@example.com", PASSWORD, role="ADMIN", is_staff=True)
        self.grant_catalog_perms(other, "view_product", "add_product")
        c = Client()
        c.force_login(other)
        self.assertEqual(c.get(reverse("api_import_detail", args=[batch["id"]])).status_code, 404)
        self.assertEqual(c.get(reverse("api_import_errors", args=[batch["id"]])).status_code, 404)
